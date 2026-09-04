from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"E:\mypaper\Cover_Letter_IEEE_TSC.docx"

BLUE = RGBColor(31, 78, 121)
GRAY = RGBColor(89, 89, 89)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_format(paragraph, before=0, after=6, line=1.10, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True


def add_text_paragraph(doc, text="", before=0, after=6, line=1.10, align=WD_ALIGN_PARAGRAPH.LEFT,
                       size=11, bold=False, italic=False, color=BLACK):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, line=line, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bottom_border(paragraph, color="B4C7E7", size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.80)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in [
    ("Title", 20, BLUE, 0, 6),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

# Restrained institutional masthead (memo_masthead pattern adapted for a journal letter).
masthead = doc.add_paragraph()
set_paragraph_format(masthead, after=2, line=1.0)
set_run_font(masthead.add_run("SHENZHEN UNIVERSITY"), size=16, bold=True, color=BLUE)

contact = doc.add_paragraph()
set_paragraph_format(contact, after=9, line=1.0)
set_run_font(
    contact.add_run("Daquan Feng, Corresponding Author  |  Shenzhen, Guangdong, China  |  fdquan@szu.edu.cn"),
    size=9.5,
    color=GRAY,
)
add_bottom_border(contact)

add_text_paragraph(doc, "September 3, 2026", before=3, after=11)

recipient = doc.add_paragraph()
set_paragraph_format(recipient, after=11, line=1.08)
set_run_font(recipient.add_run("Editor-in-Chief\n"), bold=True)
set_run_font(recipient.add_run("IEEE Transactions on Services Computing"), italic=True)

add_text_paragraph(doc, "Dear Editor-in-Chief:", after=9)

p = doc.add_paragraph()
set_paragraph_format(p, after=8)
set_run_font(p.add_run("On behalf of my coauthors, I am pleased to submit our manuscript, \""))
set_run_font(
    p.add_run("CEE-Split: Adaptive Slicing and Scheduling for Cloud-Edge-End Collaborative Rendering"),
    italic=True,
)
set_run_font(p.add_run("\", for consideration as a Regular Paper in "))
set_run_font(p.add_run("IEEE Transactions on Services Computing"), italic=True)
set_run_font(p.add_run("."))

add_text_paragraph(
    doc,
    "Interactive extended-reality applications require high visual quality and low latency, yet standalone devices remain constrained by computation, thermal, and battery limits. The submitted work addresses this problem as an adaptive service-orchestration task across heterogeneous end, edge, and cloud resources. CEE-Split jointly considers task decomposition, codec-compatible transmission and recomposition, and online resource scheduling under time-varying rendering workloads and network conditions.",
    after=8,
)

add_text_paragraph(
    doc,
    "The manuscript makes three main contributions. First, Load-Balanced Adaptive Slicing partitions visible scene geometry according to cumulative triangle workload, producing approximately balanced scheduling units without splitting individual Renderers. Second, LumaDepth RGB-D Packing carries logarithmic eye-space depth through the luma component of a single coded frame, supporting depth-aware terminal composition after lossy video transmission. Third, a PPO-based scheduler jointly selects the rendering node, rendering resolution, and transmission rate for each layer using observable workload and system telemetry. Trace-driven experiments and a three-node WebRTC prototype evaluate the complete framework. In the trace-driven study, the proposed slicing method reduces the mean layer-workload coefficient of variation from 0.445 to 0.039; compared with a training-optimized fixed configuration, the complete system reduces mean end-to-end latency from 57.84 to 53.89 ms and estimated terminal energy from 89.03 to 77.40 mJ, with a 0.54-point decrease in expected VMAF.",
    after=8,
)

add_text_paragraph(
    doc,
    "We believe the manuscript is well aligned with the journal because it studies optimization, service placement, and adaptive resource management across the cloud-edge-end continuum, while also validating the proposed methods in a working distributed rendering prototype.",
    after=8,
)

add_text_paragraph(
    doc,
    "We confirm that this manuscript is original, has not been published previously, and is not under consideration for publication elsewhere. All authors have approved the manuscript and its submission to the journal. The authors declare that they have no conflicts of interest.",
    after=8,
)

add_text_paragraph(doc, "Thank you for your consideration. We look forward to hearing from you.", after=12)

add_text_paragraph(doc, "Sincerely,", after=12)
add_text_paragraph(doc, "Daquan Feng", after=0, bold=True)
add_text_paragraph(doc, "Corresponding Author", after=0)
add_text_paragraph(doc, "Shenzhen University", after=0)
add_text_paragraph(doc, "Email: fdquan@szu.edu.cn", after=0)

doc.core_properties.title = "Cover Letter - CEE-Split"
doc.core_properties.subject = "Submission to IEEE Transactions on Services Computing"
doc.core_properties.author = "Daquan Feng"
doc.core_properties.keywords = "Cover letter; IEEE Transactions on Services Computing; CEE-Split"
doc.core_properties.comments = "Submission cover letter"

doc.save(OUTPUT)
print(OUTPUT)
