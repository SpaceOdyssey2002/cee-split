#!/usr/bin/env python3
"""
Premium PPT v4 — 大图可读 + 精美版式
- PDF/图表 400 DPI 渲染
- 论文曲线/坐标轴图：每页 ≤2 张且尽量全宽展示
- 自绘示意图加大字号
"""
from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import generate_conference_ppt_premium as C
import generate_conference_ppt_premium_v2 as L

OUT = C.BASE / "CEE_Split_Conference_Talk_Premium_v4.pptx"
E = C.EVAL
A = C.ASSETS
CHART_DPI = 400  # 坐标轴文字可读


# ── 高清示意图（更大字号）────────────────────────────────────────

def _mpl_setup():
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 11,
    })


def gen_hd_motivation():
    out = A / "ppt_v4_motivation.png"
    _mpl_setup()
    fig, ax = plt.subplots(figsize=(14, 3.8), dpi=220, facecolor="white")
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 3.8)
    ax.axis("off")
    boxes = [
        (0.4, 1.1, 2.5, 1.5, "#EBF4FF", "#2B6CB0", "XR 终端\n算力/续航受限"),
        (3.2, 1.1, 2.5, 1.5, "#ECFDF5", "#059669", "云-边-端\n协同渲染"),
        (6.0, 1.1, 2.5, 1.5, "#FEF3C7", "#B7791A", "网络波动\n负载动态"),
        (8.8, 1.1, 2.5, 1.5, "#FEE2E2", "#DC2626", "端到端\n< 100 ms"),
        (11.6, 1.1, 2.0, 1.5, "#F1F5F9", "#0F244A", "沉浸\n体验"),
    ]
    for x, y, w, h, fc, ec, txt in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.1",
                                    facecolor=fc, edgecolor=ec, linewidth=1.8))
        ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", fontsize=12, color="#1E293B")
    for i in range(4):
        ax.annotate("", xy=(boxes[i + 1][0], 1.85), xytext=(boxes[i][0] + boxes[i][2], 1.85),
                    arrowprops=dict(arrowstyle="-|>", color="#64748B", lw=1.4))
    ax.text(7, 3.35, "核心矛盾：高画质需求  vs.  严格 100ms 实时约束", ha="center", fontsize=14,
            fontweight="bold", color="#0F244A")
    fig.savefig(out, bbox_inches="tight", pad_inches=0.15, facecolor="white")
    plt.close(fig)
    return out


def gen_hd_bottlenecks():
    out = A / "ppt_v4_bottlenecks.png"
    _mpl_setup()
    fig, axes = plt.subplots(1, 3, figsize=(14, 3.4), dpi=220, facecolor="white")
    items = [
        ("#FEE2E2", "#DC2626", "静态切分", "几何负载不均\nStraggler 拖尾"),
        ("#FEF3C7", "#B7791A", "有损传输", "色度键失效\n遮挡合成错误"),
        ("#EBF4FF", "#2B6CB0", "启发式调度", "无法联合适应\n时延/画质失衡"),
    ]
    for ax, (fc, ec, title, body) in zip(axes, items):
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis("off")
        ax.add_patch(FancyBboxPatch((0.05, 0.08), 0.9, 0.84, boxstyle="round,pad=0.02,rounding_size=0.08",
                                    facecolor=fc, edgecolor=ec, linewidth=2.2))
        ax.text(0.5, 0.78, title, ha="center", fontsize=14, fontweight="bold", color=ec)
        ax.text(0.5, 0.4, body, ha="center", va="center", fontsize=12, color="#334155", linespacing=1.45)
    fig.suptitle("云-边-端协同渲染的三类关键瓶颈", fontsize=16, fontweight="bold", y=1.02, color="#0F244A")
    fig.savefig(out, bbox_inches="tight", pad_inches=0.18, facecolor="white")
    plt.close(fig)
    return out


def gen_hd_reward():
    out = A / "ppt_v4_reward.png"
    _mpl_setup()
    fig, ax = plt.subplots(figsize=(12, 4.2), dpi=220, facecolor="white")
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4.2)
    ax.axis("off")
    for x, fc, txt in [(0.3, "#ECFDF5", "r_quality\n感知加权 VMAF²"),
                       (4.2, "#FEE2E2", "r_latency\n软屏障惩罚"),
                       (8.1, "#FEF3C7", "r_switch\n抑制 Ping-Pong")]:
        ax.add_patch(FancyBboxPatch((x, 1.5), 3.4, 1.8, boxstyle="round,pad=0.02,rounding_size=0.1",
                                    facecolor=fc, edgecolor="#94A3B8", linewidth=1.2))
        ax.text(x + 1.7, 2.4, txt, ha="center", va="center", fontsize=12, color="#1E293B")
    ax.add_patch(FancyBboxPatch((2.8, 0.2), 6.4, 0.95, boxstyle="round,pad=0.02,rounding_size=0.08",
                                facecolor="#EBF4FF", edgecolor="#2B6CB0", linewidth=2))
    ax.text(6, 0.67, "R_t = r_quality^eff + r_latency − r_switch", ha="center", fontsize=14,
            fontweight="bold", color="#0F244A")
    fig.savefig(out, bbox_inches="tight", pad_inches=0.12, facecolor="white")
    plt.close(fig)
    return out


def ri(*paths: str | Path) -> Path | None:
    """Resolve chart/photo at high DPI."""
    return C.resolve_image(*paths, dpi=CHART_DPI)


# ── 版式：大图优先 ────────────────────────────────────────────────

def slide_open(prs, title_en: str, title_zh: str = "", tag: str = "", section: str = ""):
    slide = C.new_content(prs, title_en, title_zh, tag)
    if section:
        chip = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.27), Inches(2.35), Inches(0.3))
        chip.fill.solid()
        chip.fill.fore_color.rgb = C.C_TEAL
        chip.line.fill.background()
        t = slide.shapes.add_textbox(Inches(0.5), Inches(1.29), Inches(2.35), Inches(0.28))
        t.text_frame.paragraphs[0].text = section
        t.text_frame.paragraphs[0].font.size = Pt(10)
        t.text_frame.paragraphs[0].font.bold = True
        t.text_frame.paragraphs[0].font.color.rgb = C.C_WHITE
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return slide


def fig_label(slide, box: L.Box, text: str):
    if not text:
        return 0.24
    tb = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(0.22))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C.C_PRIMARY
    return 0.24


def place_fig(slide, path: Path | None, box: L.Box, caption: str = "", cap_pt: int = 10, fill_box: bool = True):
    """Place image as large as possible inside box."""
    cap_h = 0.42 if caption else 0
    label_h = 0.0
    inner_top = box.top
    inner = L.Box(box.left, inner_top, box.width, box.height - cap_h)

    if path and path.exists():
        iw, ih = L._img_size(path, inner.width - 0.1, inner.height - 0.1)
        if fill_box:
            # 尽量放大：若未沾满高度，尝试按高度放大（不超过宽度）
            scale = min(inner.width / iw, inner.height / ih)
            iw, ih = iw * scale, ih * scale
        il = inner.left + (inner.width - iw) / 2
        it = inner.top + (inner.height - ih) / 2
        C.add_card(slide, il - 0.06, it - 0.06, iw + 0.12, ih + 0.12)
        slide.shapes.add_picture(str(path), Inches(il), Inches(it), Inches(iw), Inches(ih))
    if caption:
        cb = slide.shapes.add_textbox(Inches(box.left), Inches(box.top + box.height - cap_h + 0.02),
                                      Inches(box.width), Inches(cap_h))
        cp = cb.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(cap_pt)
        cp.font.italic = True
        cp.font.color.rgb = C.C_MUTED
        cp.alignment = PP_ALIGN.CENTER


def hero(slide, path: Path | None, label: str, caption: str = ""):
    b = L.Box.full()
    lh = fig_label(slide, b, label)
    place_fig(slide, path, L.Box(b.left, b.top + lh, b.width, b.height - lh - 0.05), caption, cap_pt=11)


def duo(slide, items: list[tuple[Path | None, str, str]], box: L.Box | None = None):
    """Two large figures side by side (for PDF charts)."""
    b = box or L.Box.full()
    gap = 0.2
    w = (b.width - gap) / 2
    for i, (path, label, cap) in enumerate(items[:2]):
        cell = L.Box(b.left + i * (w + gap), b.top, w, b.height)
        lh = fig_label(slide, cell, label)
        place_fig(slide, path, L.Box(cell.left, cell.top + lh, cell.width, cell.height - lh), cap, fill_box=True)


def side_bullets_figure(slide, bullets: list[str], path: Path | None, fig_label_text: str,
                        title: str = "要点", text_w: float = 4.0):
    b = L.Box.full()
    text_box = L.Box(b.left, b.top, text_w, b.height)
    fig_box = L.Box(b.left + text_w + 0.15, b.top, b.width - text_w - 0.15, b.height)
    L.bullets_card(slide, text_box, bullets, title=title, size=13)
    lh = fig_label(slide, fig_box, fig_label_text)
    place_fig(slide, path, L.Box(fig_box.left, fig_box.top + lh, fig_box.width, fig_box.height - lh), fill_box=True)


def mini_table(slide, box: L.Box, headers, rows, font_pt=10):
    nrows, ncols = len(rows) + 1, len(headers)
    tbl = slide.shapes.add_table(nrows, ncols, Inches(box.left), Inches(box.top),
                                 Inches(box.width), Inches(box.height)).table
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C.C_NAVY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(font_pt)
            p.font.color.rgb = C.C_WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_pt)
                if "LBAS" in val or "★" in val:
                    p.font.bold = True
                    p.font.color.rgb = C.C_TEAL
                if "✗" in val:
                    p.font.color.rgb = C.C_RED
                if "✓" in val:
                    p.font.color.rgb = C.C_GREEN


def build() -> Presentation:
    mot = gen_hd_motivation()
    bot = gen_hd_bottlenecks()
    rew = gen_hd_reward()

    prs = Presentation()
    prs.slide_width = C.SLIDE_W
    prs.slide_height = C.SLIDE_H
    pg = [0]

    def finish(slide, zh: str):
        pg[0] += 1
        C.add_footer(slide, pg[0], 99)
        C.note(slide, zh)

    # ── 1 Title ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    g1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, C.SLIDE_W, Inches(3.0))
    g1.fill.solid()
    g1.fill.fore_color.rgb = C.C_NAVY
    g1.line.fill.background()
    g2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(2.55), C.SLIDE_W, Inches(0.55))
    g2.fill.solid()
    g2.fill.fore_color.rgb = C.C_ACCENT
    g2.line.fill.background()
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(3.05), C.SLIDE_W, Inches(0.07)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = C.C_TEAL
    slide.shapes[-1].line.fill.background()
    tf = slide.shapes.add_textbox(Inches(0.55), Inches(0.65), Inches(12.2), Inches(1.85)).text_frame
    tf.paragraphs[0].text = "Real-Time Cloud-Edge-End Collaborative Rendering"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = C.C_WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "via Deep Reinforcement Learning"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
    p2.alignment = PP_ALIGN.CENTER
    p3 = tf.add_paragraph()
    p3.text = "CEE Split  ·  基于深度强化学习的云-边-端协同 XR 渲染"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    p3.alignment = PP_ALIGN.CENTER
    for i, (a, b) in enumerate([
        ("LBAS", "负载均衡切分"), ("RGB-D", "深度合成"), ("PPO", "智能调度"), ("Eval", "仿真+原型")]):
        x = 0.7 + i * 3.05
        c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.55), Inches(2.85), Inches(1.05))
        c.fill.solid()
        c.fill.fore_color.rgb = C.C_CARD
        c.line.color.rgb = C.C_ACCENT
        t = slide.shapes.add_textbox(Inches(x), Inches(3.68), Inches(2.85), Inches(0.85))
        t.text_frame.paragraphs[0].text = f"{a}\n{b}"
        t.text_frame.paragraphs[0].font.size = Pt(13)
        t.text_frame.paragraphs[0].font.bold = True
        t.text_frame.paragraphs[0].font.color.rgb = C.C_PRIMARY
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    finish(slide, """【开场 · 约 45 秒】
各位老师、同学，大家好。今天汇报《基于深度强化学习的实时云-边-端协同渲染》，框架名 CEE Split。
一句话：在严格 100ms 约束下，用 LBAS 做几何均衡切分、RGB-D 保证合成正确、PPO 做联合调度。
接下来按动机、方法、实验、原型、总结展开，约 20 分钟。""")

    # 2 Agenda
    slide = slide_open(prs, "Talk Roadmap", "报告路线 (~20 min)", tag="Agenda")
    mini_table(slide, L.Box.full(), ["Part", "Time", "Content"], [
        ["1", "3 min", "Motivation & bottlenecks"],
        ["2", "6 min", "CEE Split: LBAS, RGB-D, DRL"],
        ["3", "8 min", "Evaluation & prototype"],
        ["4", "2 min", "Conclusion & Q&A"],
    ], font_pt=13)
    finish(slide, "【议程 15s】按时间控制节奏。")

    # 3 Motivation pipeline
    slide = slide_open(prs, "Motivation: Immersion vs. Latency vs. Compute", "动机：沉浸感-时延-算力", section="§1")
    hero(slide, mot, "XR 应用的核心矛盾链", "高画质需求 · 移动终端受限 · 100ms 硬约束")
    finish(slide, """【动机 1 · 90s】4K/GI 需求、头显受限、必须卸载、100ms 来自 3GPP、超时会晕动。""")

    # 4 Bottlenecks
    slide = slide_open(prs, "Three System Bottlenecks", "三类系统瓶颈", section="§1")
    hero(slide, bot, "现有方案为何无法同时满足画质与时延")
    finish(slide, """【动机 2 · 90s】静态切分 Straggler；VP8 色度键失败；启发式无法联合优化。""")

    # 5 Framework
    slide = slide_open(prs, "CEE Split Framework", "CEE Split 框架", section="§2", tag="Core")
    side_bullets_figure(slide, [
        "LBAS：几何负载均衡切分 Near/Mid/Far",
        "RGB-D：luma 传深度 + Z-Test 合成",
        "DRL：PPO 满足 100ms 最大化 QoE",
        "双环：端侧推理 + 云端训练",
    ], ri(A / "fig_macro_architecture.eps", A / "fig_macro_architecture-eps-converted-to.pdf"),
                      "云-边-端宏观架构", text_w=4.2)
    finish(slide, "【框架 75s】数据面+控制面协同。")

    # 6 Contributions
    slide = slide_open(prs, "Contributions", "主要贡献", section="§2")
    b = L.Box.full()
    for i, (cid, en, zh) in enumerate([
        ("C1", "LBAS + RGB-D", "可切分、可重组的数据面"),
        ("C2", "MDP + DRL", "10D 状态 / 软屏障奖励"),
        ("C3", "Sim + Prototype", "轨迹仿真 + WebRTC 三节点"),
    ]):
        c = b.col(i, 3)
        C.add_card(slide, c.left, c.top, c.width, c.height, fill=C.C_SOFT_BLUE)
        t = slide.shapes.add_textbox(Inches(c.left + 0.15), Inches(c.top + 0.2), Inches(c.width - 0.25), Inches(c.height - 0.3))
        t.text_frame.paragraphs[0].text = f"{cid}\n{en}\n{zh}"
        t.text_frame.paragraphs[0].font.size = Pt(14)
        t.text_frame.paragraphs[0].font.bold = True
        t.text_frame.paragraphs[0].font.color.rgb = C.C_PRIMARY
    finish(slide, "【贡献 60s】三点逐一强调。")

    # 7 System tiers
    slide = slide_open(prs, "Three-Tier Architecture", "三层架构", section="§3")
    side_bullets_figure(slide, [
        "End：LBAS + PPO 推理 + 合成/ATW",
        "Edge：低时延渲染 + 热缓存",
        "Cloud：重负载 + 全局 PPO 训练",
        "控制上行 / 视频下行",
    ], ri(A / "fig_macro_architecture.eps", A / "fig_macro_architecture-eps-converted-to.pdf"),
                      "架构详解", text_w=3.8)
    finish(slide, "【架构 60s】")

    # 8 Pipeline — 全宽大图
    slide = slide_open(prs, "Parallel Rendering Pipeline & Dual Loops", "并行流水线与双环控制", section="§3")
    hero(slide, ri(A / "fig_pipeline_timing.eps", A / "fig_pipeline_timing-eps-converted-to.pdf"),
         "分布式渲染时序（并行分支 + 双环训练）",
         "L_final = max(L_near, L_mid, L_far) + T_comp")
    finish(slide, "【流水线 75s】强调 max 分支时延模型。")

    # 9 LBAS
    slide = slide_open(prs, "LBAS: Load-Balanced Adaptive Slicing", "LBAS 负载均衡自适应切分", section="§4")
    side_bullets_figure(slide, [
        "w_k = 三角面数，沿深度累积 L(z)",
        "d_near, d_mid：各约 1/3 几何量",
        "指数平滑避免平面抖动",
        "LBAS 均衡，DRL 选节点/画质",
    ], ri(A / "fig_lbas_principle.eps", A / "fig_lbas_principle-eps-converted-to.pdf"),
                      "LBAS 几何切分原理", text_w=4.0)
    finish(slide, "【LBAS 90s】")

    # 10 LBAS surge — 每张图一页或半页
    slide = slide_open(prs, "LBAS Under Workload Surge (1/2)", "负载激增实验 (1/2)", section="§4")
    duo(slide, [
        (ri(E / "Fig6a_Fixed_Depth.pdf"), "(a) 固定深度切分", "三角形分布失衡"),
        (ri(E / "Fig6b_LBAS.pdf"), "(b) LBAS 重平衡", "负载沿深度均衡"),
    ])
    finish(slide, "【LBAS 实验 30s】对比固定深度。")

    slide = slide_open(prs, "LBAS Under Workload Surge (2/2)", "负载激增实验 (2/2) — 时延", section="§4")
    hero(slide, ri(E / "Fig6c_Overload_Latency.pdf"), "(c) 端到端时延：仅 LBAS 守住 100ms",
         "帧 30–70 几何激增区间")
    finish(slide, """【LBAS 时延 · 约 45 秒】
本页全宽大图，请指着 30–70 帧区间：固定深度方案时延冲破 100ms，LBAS 保持平稳。
这是回应「为何不用固定前景/背景」的关键实验页。""")

    # 11–13 RGB-D
    slide = slide_open(prs, "RGB-D Packing Pipeline", "RGB-D 打包数据通路", section="§5")
    hero(slide, ri(A / "fig_rgbd_layout_preview_v2.png", A / "fig_rgbd_layout_preview_v2.pdf"),
         "空间复用 RGB-D：上半 RGB / 下半深度(luma)")
    finish(slide, "【RGB-D 60s】")

    slide = slide_open(prs, "Compositing: Chroma-Key vs. Z-Test", "合成：色度键 vs 深度 Z-Test", section="§5")
    hero(slide, C.resolve_image(A / "fig_rgbd_packing.jpg", dpi=200),
         "真实 VP8 链路：色度键失败（左）vs 深度合成（右）")
    finish(slide, "【合成 45s】")

    slide = slide_open(prs, "Perceptual Asymmetry by Depth Layer", "深度层感知不对称", section="§5")
    b = L.Box.full()
    duo(slide, [
        (C.resolve_image(A / "fig_visual_a.png"), "(a) Baseline", ""),
        (C.resolve_image(A / "fig_visual_b.png"), "(b) Near degraded", ""),
    ], L.Box(b.left, b.top, b.width, b.height / 2 - 0.08))
    duo(slide, [
        (C.resolve_image(A / "fig_visual_c.png"), "(c) Mid degraded", ""),
        (C.resolve_image(A / "fig_visual_d.png"), "(d) Far degraded", ""),
    ], L.Box(b.left, b.top + b.height / 2 + 0.08, b.width, b.height / 2 - 0.08))
    finish(slide, "【感知 45s】ω=[0.5,0.35,0.15]。")

    # 14–16 DRL
    slide = slide_open(prs, "MDP Formulation", "MDP 建模", section="§6", tag="DRL")
    hero(slide, ri(A / "fig_drl_mdp.png", A / "fig_drl_mdp.eps"), "状态-动作-环境交互")
    finish(slide, "【MDP 60s】10D 状态，6D Multi-Discrete 动作。")

    slide = slide_open(prs, "Composite Reward Function", "复合奖励函数", section="§6")
    side_bullets_figure(slide, [
        "r_quality：VMAF² + 85/95 bonus",
        "违约：r_quality 减半",
        "r_latency：软屏障多项式",
        "r_switch：抑制 Ping-Pong",
    ], rew, "奖励结构", text_w=4.5)
    finish(slide, "【奖励 60s】")

    slide = slide_open(prs, "Training Convergence", "PPO 训练收敛", section="§6")
    hero(slide, ri(A / "fig_training_curve.png"), "端到端时延收敛至 100ms 以下")
    finish(slide, "【训练 40s】")

    slide = slide_open(prs, "QoE Proxy Model", "画质代理模型", section="§6")
    duo(slide, [
        (ri(A / "fig_loss_curve.png"), "代理网络训练损失", ""),
        (ri(A / "fig_vmaf_scatter.png"), "VMAF 预测 (MAE≈2.6)", ""),
    ])
    finish(slide, "【代理 30s】<1ms 前向。")

    # 17–22 Evaluation — 每张论文曲线单独大图
    slide = slide_open(prs, "Evaluation Setup", "实验设置", section="§7")
    L.bullets_card(slide, L.Box.full(), [
        "Unity + Gymnasium 轨迹驱动仿真",
        "基线：End Only / Cloud Only / Heuristic",
        "五场景：Light · Excellent · Fluctuating · Congested · Heavy",
        "PPO 2M steps；域随机化 c_end, c_edge, c_cloud",
        "指标：VMAF(代理)、E2E 时延、终端能耗",
    ], title="仿真配置", size=14)
    finish(slide, "【设置 45s】")

    slide = slide_open(prs, "Results: Latency Across Scenarios", "结果：各场景时延", section="§7")
    hero(slide, ri(E / "Fig1a_Latency.pdf"), "平均时延对比（log scale）")
    finish(slide, "【Fig1a 30s】")

    slide = slide_open(prs, "Results: Visual Quality", "结果：视觉质量", section="§7")
    hero(slide, ri(E / "Fig1b_Quality.pdf"), "平均 VMAF 对比")
    finish(slide, "【Fig1b 30s】")

    slide = slide_open(prs, "Results: Robustness", "结果：鲁棒性", section="§7")
    hero(slide, ri(E / "Fig1c_Robustness.pdf"), "画质鲁棒性")
    finish(slide, "【Fig1c 20s】")

    slide = slide_open(prs, "Quantitative Summary", "量化结果摘要", section="§7", tag="Key")
    b = L.Box.full()
    mini_table(slide, b, ["场景", "策略", "VMAF", "时延(ms)"], [
        ["Congested", "Heuristic", "67.6", "125.7✗"],
        ["Congested", "LBAS-DRL★", "79.0", "90.5✓"],
        ["Heavy", "End Only", "87.6", "167.1✗"],
        ["Heavy", "LBAS-DRL★", "82.5", "56.1✓"],
        ["Light", "Cloud Only", "95.4", "141.1✗"],
        ["Light", "LBAS-DRL★", "97.4", "37.3✓"],
    ], font_pt=11)
    finish(slide, """【核心表格 · 约 90 秒】
请放慢语速读三组对比。拥塞与重负载时仅 LBAS-DRL 守住 100ms；轻载 VMAF 97.4 超过 Cloud 上限 95.4。
这是全文最核心的量化结论，务必让听众记住。""")

    slide = slide_open(prs, "Pareto Frontier", "帕累托前沿", section="§7")
    hero(slide, ri(E / "Fig3_Pareto_Frontier.pdf"), "时延-画质帕累托支配关系")
    finish(slide, "【Pareto 40s】")

    slide = slide_open(prs, "Tail Latency CDF", "尾部时延 CDF", section="§7")
    hero(slide, ri(E / "Fig2_Latency_CDF.pdf"), "100ms 前收敛至 100%")
    finish(slide, "【CDF 40s】")

    slide = slide_open(prs, "Dynamic Response to Bandwidth Drop", "带宽骤降的动态响应", section="§7")
    hero(slide, ri(E / "Fig4_Dynamic_Response.pdf"), "帧 30 带宽骤降：立即降质保时延")
    finish(slide, "【动态 40s】")

    slide = slide_open(prs, "Learned Offloading Strategy", "学到的卸载策略分布", section="§7")
    hero(slide, ri(E / "Fig5_Strategy_Distribution.pdf"), "不同场景下的节点选择")
    finish(slide, "【策略 40s】")

    slide = slide_open(prs, "Layer-wise Quality & Node Assignment", "分层画质与节点分配", section="§7")
    duo(slide, [
        (ri(E / "Fig7a_Layer_VMAF.pdf"), "各层 VMAF", ""),
        (ri(E / "Fig7b_Layer_Node.pdf"), "各层渲染节点", ""),
    ])
    finish(slide, "【分层 40s】保近压远。")

    slide = slide_open(prs, "Energy & Zero-Shot Generalization", "能耗与零样本泛化", section="§7")
    duo(slide, [
        (ri(E / "Fig8_Energy_Analysis.pdf"), "终端能耗", ""),
        (ri(E / "Fig10_Domain_Randomization.pdf"), "跨硬件泛化", ""),
    ])
    finish(slide, "【能耗/泛化 30s】")

    slide = slide_open(prs, "Reward Ablation Study", "奖励消融", section="§7")
    duo(slide, [
        (ri(E / "Fig9a_Ablation_CDF.pdf"), "去掉软屏障", ""),
        (ri(E / "Fig9b_Ablation_Switch.pdf"), "去掉切换惩罚", ""),
    ])
    finish(slide, "【消融 40s】")

    # Prototype
    slide = slide_open(prs, "Prototype Testbed", "物理原型测试床", section="§8")
    b = L.Box.full()
    top, bottom = b.split_rows(2, gap=0.12)
    for i, txt in enumerate(["Cloud · RTX 4060", "Edge · RTX 2070", "End · SD 7 Gen 3"]):
        c = top.col(i, 3)
        C.add_card(slide, c.left, c.top, c.width, c.height)
        t = slide.shapes.add_textbox(Inches(c.left), Inches(c.top + 0.15), Inches(c.width), Inches(c.height - 0.2))
        t.text_frame.paragraphs[0].text = txt + "\nTailscale · VP8/WebRTC"
        t.text_frame.paragraphs[0].font.size = Pt(12)
        t.text_frame.paragraphs[0].font.bold = True
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    duo(slide, [
        (C.resolve_image(A / "fig_courtyard_total.png", dpi=220), "Courtyard 场景", ""),
        (C.resolve_image(A / "fig_courtyard_split.png", dpi=220), "LBAS 三层切分", ""),
    ], bottom)
    finish(slide, "【原型设置 45s】")

    slide = slide_open(prs, "Codec & Depth Map Quality", "编解码与深度图", section="§8")
    duo(slide, [
        (C.resolve_image(A / "fig_gt.png", dpi=220), "Ground Truth", ""),
        (C.resolve_image(A / "fig_vp8_10k.png", dpi=220), "VP8 @10Mbps", ""),
    ])
    finish(slide, "【Codec 30s】")

    slide = slide_open(prs, "Prototype End-to-End Performance", "原型端到端性能", section="§8")
    duo(slide, [
        (ri(A / "Fig1_Performance_Comparison.pdf"), "VMAF–时延", ""),
        (ri(A / "Fig2_Latency_Breakdown.pdf"), "时延分解", ""),
    ])
    b = L.Box(L.X0, L.Y1 - 1.35, L.CONTENT_W, 1.25)
    L.stat_row(slide, [
        ("36ms", "LBAS-DRL\nVMAF 85.17", C.C_GREEN),
        ("98.1ms", "Cloud Only", C.C_ACCENT),
        ("122.5ms", "Heuristic B ✗", C.C_RED),
    ], b)
    finish(slide, "【原型结果 60s】")

    # Related + Conclusion
    slide = slide_open(prs, "Positioning vs. Prior Work", "与已有工作对比", section="§9")
    mini_table(slide, L.Box.full(), ["方向", "局限", "CEE Split"], [
        ["整帧云游戏", "带宽/时延敏感", "多层并行+LBAS"],
        ["固定深度切分", "Straggler", "自适应平面"],
        ["DRL for ABR", "非 3D 管线", "渲染-网络 MDP"],
        ["色度键合成", "VP8 失效", "RGB-D+Z-Test"],
    ], font_pt=11)
    finish(slide, "【相关 45s】")

    slide = slide_open(prs, "Takeaways & Future Work", "总结与展望", section="§10")
    b = L.Box.full(bottom_pad=1.4)
    top, bot = b.split_rows(2)
    for i, (title, items) in enumerate([
        ("贡献", ["LBAS+RGB-D+DRL 端到端", "全场景<100ms", "仿真+原型验证"]),
        ("局限", ["仿真为主", "单用户", "资产预同步"]),
        ("未来", ["MARL 多用户", "MEC 大规模实测", "NeRF/3DGS"]),
    ]):
        c = top.col(i, 3)
        L.bullets_card(slide, c, items, title=title, size=12)
    L.stat_row(slide, [("100%", "时延合规", C.C_GREEN), ("97.4", "Light VMAF", C.C_ACCENT), ("36ms", "Prototype", C.C_TEAL)], bot)
    finish(slide, "【总结 90s】")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, C.SLIDE_W, C.SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C.C_NAVY
    bg.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2))
    tb.text_frame.paragraphs[0].text = "Thank You  ·  谢谢"
    tb.text_frame.paragraphs[0].font.size = Pt(50)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = C.C_WHITE
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    p2 = tb.text_frame.add_paragraph()
    p2.text = "Questions & Discussion"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    p2.alignment = PP_ALIGN.CENTER
    finish(slide, "【Q&A】")

    n = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if " / 99" in p.text:
                        p.text = f"{i + 1} / {n}"
    return prs


def main():
    print("Generating premium v4 (large readable figures)...")
    prs = build()
    prs.save(str(OUT))
    print(f"Saved: {OUT}")
    print(f"Slides: {len(prs.slides)}")
    print("Charts rendered at 400 DPI — axis labels should be readable in slideshow.")


if __name__ == "__main__":
    main()
